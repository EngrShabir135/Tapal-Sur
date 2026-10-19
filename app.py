"""
Tapal DD 900gm Collectible CP - Streamlit daily updater.
Run with:  streamlit run app.py
Requirements:  pip install streamlit pandas openpyxl
"""
import streamlit as st
import pandas as pd

import re
import io
import csv
import copy
import difflib
import datetime as dt

import openpyxl
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI
from openpyxl.utils.cell import range_boundaries
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.workbook.properties import CalcProperties

MON = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
       "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
CS_DAY_RE = re.compile(r"^(\d{1,2})-([A-Za-z]{3})$")                       # "15-Oct"
DD_DAY_RE = re.compile(r"^Sellout \n\((\d{1,2})(?:st|nd|rd|th) ([A-Za-z]{3}) (\d{4})\)$")
DD_STATUS_PREFIX = "DD 900gm CP Stock Status"
REF_RE = re.compile(r"(?<![A-Za-z0-9_.\[!'])(\$?)([A-Z]{1,3})(\$?)(\d+)(?![A-Za-z0-9_(\[!])")
UPDATED_TILL_RE = re.compile(r"updated till \d{1,2} [A-Za-z]{3} \d{4}")
FUZZY_CUTOFF = 0.92
STOCK_FINISHED = "Stock finished"


# ----------------------------------------------------------------- helpers
def norm(s):
    return re.sub(r"\s+", " ", str(s or "")).strip().lower()


def clean(s):
    return re.sub(r"\s+", " ", str(s or "")).strip()


def ordinal(n):
    suffix = "th" if 11 <= n % 100 <= 13 else {1: "st", 2: "nd", 3: "rd"}.get(n % 10, "th")
    return f"{n}{suffix}"


def day_key(d):
    return (d.day, MON[d.month - 1])


def cs_label(d):
    return f"{d.day}-{MON[d.month - 1]}"


def dd_label(d):
    return f"Sellout \n({ordinal(d.day)} {MON[d.month - 1]} {d.year})"


def dd_status_label(d):
    return f"DD 900gm CP Stock Status\n({ordinal(d.day)} {MON[d.month - 1]} {d.year})"


def dd_remarks_label(d):
    return f"Remarks ({d.day:02d}/{d.month:02d}/{d.year})"


def to_num(s):
    s = (s or "").strip()
    if re.fullmatch(r"-?\d+(\.\d+)?", s):
        f = float(s)
        return int(f) if f.is_integer() else f
    return None


def stock_status(q1):
    """Return (status or None, is_out_of_stock_answer) from the Q1 CP answer."""
    s = (q1 or "").strip()
    if not s:
        return None, False
    low = s.lower()
    if low == "out of stock":
        return "Not Available", True
    if low in ("not available", "n/a", "na"):
        return "Not Available", False
    v = to_num(s)
    if v is None:
        return None, False
    return ("Available" if v > 0 else "Not Available"), False


# ----------------------------------------------------------------- survey
def _parse_ts(time_s, d):
    """Survey times are not zero-padded (e.g. '20:9:6'), so parse the parts."""
    try:
        h, m, s = [int(x) for x in (time_s or "").strip().split(":")]
        return dt.datetime(d.year, d.month, d.day, h, m, s)
    except (ValueError, TypeError):
        return None


def read_survey(raw: bytes):
    """Parse the Tapal raw CSV into one record per submission (ObjectID)."""
    text = raw.decode("utf-8-sig", errors="replace")
    rows = list(csv.reader(io.StringIO(text)))
    header = [h.strip() for h in rows[0]]
    subs = {}
    for r in rows[1:]:
        rec = dict(zip(header, r))              # extra trailing fields are ignored
        oid = (rec.get("ObjectID") or "").strip()
        if not oid:
            continue
        store = clean(rec.get("Store"))
        date_s = (rec.get("Date") or "").strip()
        try:
            d = dt.date.fromisoformat(date_s[:10])
        except ValueError:
            continue
        if not store:
            continue
        s = subs.setdefault(oid, {"store": store, "date": d, "ts": None, "ans": {}})
        s["ts"] = _parse_ts(rec.get("Time"), d) or s["ts"]
        m = re.match(r"\s*Q(\d)\.", rec.get("Question") or "")
        if m:
            s["ans"]["Q" + m.group(1)] = (rec.get("Answer") or "").strip()
    return list(subs.values())


def pick_visits(subs):
    """One visit per (store, date). Duplicates: latest submission that has answers."""
    groups = {}
    for s in subs:
        groups.setdefault((norm(s["store"]), s["date"]), []).append(s)

    visits, dups = {}, []
    for (skey, d), lst in groups.items():
        lst.sort(key=lambda s: s["ts"] or dt.datetime.min, reverse=True)
        usable = [s for s in lst if s["ans"].get("Q1") or s["ans"].get("Q2")]
        chosen = (usable or lst)[0]
        visits[(skey, d)] = chosen
        if len(lst) > 1:
            identical = all(s["ans"] == lst[0]["ans"] for s in lst)
            dups.append({
                "Store": lst[0]["store"], "Date": d.isoformat(),
                "Submissions": len(lst),
                "Used": chosen["ts"].strftime("%H:%M:%S") if chosen["ts"] else "",
                "Identical": "Yes" if identical else "No",
            })
    return visits, dups


def match_stores(survey_keys, wb_keys):
    wb_set = set(wb_keys)
    result = {}
    for k in survey_keys:
        if k in wb_set:
            result[k] = (k, "exact")
            continue
        best = difflib.get_close_matches(k, wb_keys, n=1, cutoff=FUZZY_CUTOFF)
        result[k] = (best[0], "fuzzy") if best else (None, "unmatched")
    return result


# ----------------------------------------------------------------- sheet surgery
def shift_formula(f, threshold, n):
    def rep(m):
        c = CI(m.group(2))
        if c >= threshold:
            c += n
        return f"{m.group(1)}{L(c)}{m.group(3)}{m.group(4)}"
    return REF_RE.sub(rep, f)


def shift_rng(rng, threshold, n):
    out = []
    for part in rng.split(":"):
        m = re.match(r"([A-Z]+)(\d+)$", part)
        c = CI(m.group(1))
        if c >= threshold:
            c += n
        out.append(f"{L(c)}{m.group(2)}")
    return ":".join(out)


def find_header_row(ws):
    for r in range(1, 15):
        vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if "LMT Store" in vals:
            return r
    raise ValueError(f"Sheet '{ws.title}': could not find the 'LMT Store' header row")


def is_data_sheet(ws):
    try:
        find_header_row(ws)
        return True
    except ValueError:
        return False


def scan_headers(ws, hdr_row):
    return {c: ws.cell(hdr_row, c).value for c in range(1, ws.max_column + 1)
            if isinstance(ws.cell(hdr_row, c).value, str)}


def insert_columns(ws, threshold, n, tbl, new_names, old_last, new_last):
    """Insert n columns at `threshold` and keep formulas, table, merges, CF and widths in sync."""
    merges = [str(m) for m in ws.merged_cells.ranges]
    for m in merges:
        if range_boundaries(m)[2] >= threshold:
            ws.unmerge_cells(m)

    cfs = [([str(r) for r in cf.sqref.ranges], list(cf.rules)) for cf in ws.conditional_formatting]
    auto_ref = ws.auto_filter.ref

    ws.insert_cols(threshold, n)

    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                v = shift_formula(c.value, threshold, n)
                if old_last and new_last and old_last != new_last:
                    v = v.replace(f":[{old_last}]]", f":[{new_last}]]")
                c.value = v

    for m in merges:
        if range_boundaries(m)[2] >= threshold:
            ws.merge_cells(shift_rng(m, threshold, n))

    new_cf = ConditionalFormattingList()
    for rngs, rules in cfs:
        sq = " ".join(shift_rng(r, threshold, n) for r in rngs)
        for rule in rules:
            new_cf.add(sq, rule)
    ws.conditional_formatting = new_cf

    if auto_ref:
        ws.auto_filter.ref = shift_rng(auto_ref, threshold, n)

    dims = {}
    for k, d in list(ws.column_dimensions.items()):
        lo = d.min or CI(k)
        hi = d.max or lo
        for i in range(lo, min(hi, 200) + 1):
            if d.width:
                dims[i] = d.width
    ws.column_dimensions.clear()
    for i, w in dims.items():
        j = i + n if i >= threshold else i
        ws.column_dimensions[L(j)].width = w
    if threshold - 1 in dims:
        for i in range(n):
            ws.column_dimensions[L(threshold + i)].width = dims[threshold - 1]

    # table: extend ref, add columns after the last day column, drop stale sortState
    min_c, min_r, max_c, max_r = range_boundaries(tbl.ref)
    if max_c >= threshold:
        max_c += n
    tbl.ref = f"{L(min_c)}{min_r}:{L(max_c)}{max_r}"
    if tbl.autoFilter is not None:
        tbl.autoFilter.ref = tbl.ref
    tbl.sortState = None
    for tc in tbl.tableColumns:                      # e.g. Sum column's calculated formula
        if tc.calculatedColumnFormula is not None and old_last and new_last:
            txt = tc.calculatedColumnFormula.attr_text
            tc.calculatedColumnFormula.attr_text = txt.replace(f":[{old_last}]]", f":[{new_last}]]")

    names = [c.name for c in tbl.tableColumns]
    idx = names.index(old_last)
    tmpl = tbl.tableColumns[idx]
    next_id = max(c.id for c in tbl.tableColumns) + 1
    new_cols = []
    for i, nm in enumerate(new_names):
        tc = copy.copy(tmpl)
        tc.id = next_id + i
        tc.name = nm
        tc.calculatedColumnFormula = None
        tc.totalsRowFunction = None
        tc.totalsRowLabel = None
        new_cols.append(tc)
    tbl.tableColumns[idx + 1: idx + 1] = new_cols


def rename_table_column(tbl, old, new):
    for c in tbl.tableColumns:
        if c.name == old:
            c.name = new


# ----------------------------------------------------------------- per-sheet update
def update_sheet(ws, kind, visits_for_wb, all_dates, log):
    """kind = 'cs' (Collectible Summary) or 'dd' (DD 900gm Compile)."""
    hdr_row = find_header_row(ws)
    hdr = scan_headers(ws, hdr_row)
    store_col = next(c for c, v in hdr.items() if v == "LMT Store")
    total_col = next(c for c, v in hdr.items() if v in ("Sum", "Total Sale"))
    status_col = next(c for c, v in hdr.items()
                      if v == "Status" or (isinstance(v, str) and v.startswith(DD_STATUS_PREFIX)))
    remarks_col = next(c for c, v in hdr.items()
                       if v == "Remarks" or (isinstance(v, str) and v.startswith("Remarks (")))
    day_re = CS_DAY_RE if kind == "cs" else DD_DAY_RE

    def existing_days(h):
        out = {}
        for c, v in h.items():
            m = day_re.match(v)
            if m:
                out[(int(m.group(1)), m.group(2).title())] = c
        return out

    days = existing_days(hdr)
    old_last_col = max(days.values())
    old_last_label = hdr[old_last_col]
    new_dates = sorted({d for d in all_dates if day_key(d) not in days})
    tbl = list(ws.tables.values())[0]

    if new_dates:
        n = len(new_dates)
        new_labels = [cs_label(d) if kind == "cs" else dd_label(d) for d in new_dates]
        new_last_label = new_labels[-1]
        insert_columns(ws, total_col, n, tbl, new_labels, old_last_label, new_last_label)
        for i, lab in enumerate(new_labels):
            ws.cell(hdr_row, total_col + i).value = lab
        for r in range(hdr_row, ws.max_row + 1):
            src = ws.cell(r, total_col - 1)
            for i in range(n):
                ws.cell(r, total_col + i)._style = copy.copy(src._style)
        log.append(f"Added {n} new day column(s) to '{ws.title}': "
                   + ", ".join(d.strftime("%d-%b-%Y") for d in new_dates))
        # DD: totals row for the new day columns
        if kind == "dd":
            total_row = next((r for r in range(hdr_row + 1, ws.max_row + 1)
                              if ws.cell(r, 6).value == "Total"), None)
            if total_row:
                for i in range(n):
                    col = total_col + i
                    ws.cell(total_row, col).value = (
                        f"=SUM({L(col)}{hdr_row + 1}:{L(col)}{total_row - 1})")
        days = existing_days({c: ws.cell(hdr_row, c).value for c in range(1, ws.max_column + 1)
                              if isinstance(ws.cell(hdr_row, c).value, str)})
        hdr = scan_headers(ws, hdr_row)
        total_col = next(c for c, v in hdr.items() if v in ("Sum", "Total Sale"))
        status_col = next(c for c, v in hdr.items()
                          if v == "Status" or (isinstance(v, str) and v.startswith(DD_STATUS_PREFIX)))
        remarks_col = next(c for c, v in hdr.items()
                           if v == "Remarks" or (isinstance(v, str) and v.startswith("Remarks (")))

    # latest day in the sheet drives the status header rename (DD only)
    latest_date = max(all_dates)
    if kind == "dd":
        old_status = hdr[status_col]
        new_status = dd_status_label(latest_date)
        if old_status != new_status:
            ws.cell(hdr_row, status_col).value = new_status
            rename_table_column(tbl, old_status, new_status)
            log.append(f"Renamed stock status header to '{ws.title}' -> "
                       f"{latest_date.strftime('%d-%b-%Y')}")
        ws.cell(hdr_row, remarks_col).value = dd_remarks_label(latest_date)
        old_status_for_wb = old_status
    else:
        old_status_for_wb = None

    # write survey values
    updated_rows = 0
    status_changes = 0
    matched_keys = set()
    for r in range(hdr_row + 1, ws.max_row + 1):
        skey = norm(ws.cell(r, store_col).value)
        if not skey or skey not in visits_for_wb:
            continue
        matched_keys.add(skey)
        vis = visits_for_wb[skey]
        touched = False
        for d, v in sorted(vis.items()):
            col = days.get(day_key(d))
            if col is None:
                continue
            num = to_num(v["ans"].get("Q2", ""))
            if num is not None:
                ws.cell(r, col).value = num
                touched = True
        if touched:
            updated_rows += 1

        latest_q1 = [v for d, v in sorted(vis.items()) if v["ans"].get("Q1", "").strip()]
        if latest_q1:
            v = latest_q1[-1]
            st, oos = stock_status(v["ans"]["Q1"])
            if st:
                prev = ws.cell(r, status_col).value
                ws.cell(r, status_col).value = st
                rem = ws.cell(r, remarks_col)
                if st == "Available" and rem.value == STOCK_FINISHED:
                    rem.value = None
                elif st == "Not Available" and (oos or prev == "Available") and not rem.value:
                    rem.value = STOCK_FINISHED
                if prev != st:
                    status_changes += 1

    # "Data is updated till ..." notes
    note = f"updated till {latest_date.day} {MON[latest_date.month - 1]} {latest_date.year}"
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and UPDATED_TILL_RE.search(c.value):
                c.value = UPDATED_TILL_RE.sub(note, c.value)

    return {
        "sheet": ws.title,
        "new_dates": new_dates,
        "rows_updated": updated_rows,
        "status_changes": status_changes,
        "old_status_header": old_status_for_wb,
        "matched_keys": matched_keys,
    }


def process_workbook(raw_bytes, kind, survey_visits_by_key, all_dates, title):
    """Update one master workbook. Returns (bytes, info dict)."""
    wb = openpyxl.load_workbook(io.BytesIO(raw_bytes))
    log = []
    sheet_infos = []

    # which survey stores map to which workbook store key (per sheet, all sheets together)
    wb_keys = []
    data_sheets = [ws for ws in wb.worksheets if is_data_sheet(ws)]
    for ws in data_sheets:
        hdr_row = find_header_row(ws)
        store_col = next(c for c in range(1, ws.max_column + 1)
                         if ws.cell(hdr_row, c).value == "LMT Store")
        for r in range(hdr_row + 1, ws.max_row + 1):
            k = norm(ws.cell(r, store_col).value)
            if k:
                wb_keys.append(k)
    wb_keys = list(dict.fromkeys(wb_keys))
    mapping = match_stores(list(survey_visits_by_key.keys()), wb_keys)

    visits_for_wb = {}
    fuzzy, unmatched = [], []
    for skey, (wkey, how) in mapping.items():
        if how == "unmatched":
            unmatched.append(skey)
            continue
        if how == "fuzzy":
            fuzzy.append((skey, wkey))
        visits_for_wb.setdefault(wkey, {}).update(survey_visits_by_key[skey])

    for ws in data_sheets:
        info = update_sheet(ws, kind, visits_for_wb, all_dates, log)
        sheet_infos.append(info)

    # DD: stock status header is referenced by lookup formulas on other sheets
    if kind == "dd":
        old = sheet_infos[0]["old_status_header"]
        new = dd_status_label(max(all_dates))
        if old and old != new:
            for ws in wb.worksheets:
                for row in ws.iter_rows():
                    for c in row:
                        if isinstance(c.value, str) and c.value.startswith("=") and f"[{old}]" in c.value:
                            c.value = c.value.replace(f"[{old}]", f"[{new}]")

    # drop slicer names (slicers themselves are not preserved by openpyxl)
    for name in list(wb.defined_names.keys()):
        if name.startswith("Slicer_"):
            del wb.defined_names[name]

    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    add_log_sheet(wb, title, log, unmatched, fuzzy)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), {"sheets": sheet_infos, "fuzzy": fuzzy, "unmatched": unmatched, "log": log}


def add_log_sheet(wb, title, log, unmatched, fuzzy):
    if "Update Log" in wb.sheetnames:
        del wb["Update Log"]
    ws = wb.create_sheet("Update Log")
    ws.append([f"Update run: {title}", dt.datetime.now().strftime("%d-%b-%Y %H:%M")])
    ws.append([])
    for line in log:
        ws.append([line])
    if fuzzy:
        ws.append([])
        ws.append(["Fuzzy-matched store names (please verify):"])
        for s, w in fuzzy:
            ws.append([f"Survey: '{s}'  ->  Workbook: '{w}'"])
    if unmatched:
        ws.append([])
        ws.append(["Survey stores NOT found in this workbook (not updated):"])
        for s in unmatched:
            ws.append([f"  {s}"])
    ws.column_dimensions["A"].width = 110
    ws["A1"].font = openpyxl.styles.Font(bold=True)


# ----------------------------------------------------------------- public entry point
def run_update(survey_csv: bytes, cs_xlsx: bytes, dd_xlsx: bytes):
    subs = read_survey(survey_csv)
    if not subs:
        raise ValueError("No valid survey rows found in the CSV (check the file).")
    visits, dups = pick_visits(subs)

    by_store = {}
    for (skey, d), v in visits.items():
        by_store.setdefault(skey, {})[d] = v
    all_dates = sorted({d for (_, d) in visits})

    cs_bytes, cs_info = process_workbook(cs_xlsx, "cs", by_store, all_dates, "Collectible Summary")
    dd_bytes, dd_info = process_workbook(dd_xlsx, "dd", by_store, all_dates, "DD 900gm CP Survey")

    # per-store report for display
    report = []
    for skey, dates in sorted(by_store.items()):
        name = next(v["store"] for v in dates.values())
        q1 = [v for d, v in sorted(dates.items()) if v["ans"].get("Q1", "").strip()]
        st = stock_status(q1[-1]["ans"]["Q1"])[0] if q1 else None
        sold = []
        for d, v in sorted(dates.items()):
            num = to_num(v["ans"].get("Q2", ""))
            sold.append(f"{d.strftime('%d-%b')}: {num if num is not None else '-'}")
        report.append({
            "Store": name,
            "Days surveyed": ", ".join(d.strftime("%d-%b") for d in sorted(dates)),
            "Units sold (Q2) by day": "; ".join(sold),
            "Latest stock status": st or "-",
        })
    daily = {}
    for (skey, d), v in visits.items():
        num = to_num(v["ans"].get("Q2", ""))
        if num is not None:
            daily[d] = daily.get(d, 0) + num

    return {
        "daily_totals": daily,
        "cs_bytes": cs_bytes,
        "dd_bytes": dd_bytes,
        "dates": all_dates,
        "submissions": len(subs),
        "stores_surveyed": len(by_store),
        "duplicates": dups,
        "cs_info": cs_info,
        "dd_info": dd_info,
        "report": report,
    }


# =============================================================================
# Streamlit UI
# =============================================================================
st.set_page_config(page_title="Tapal Daily Survey Updater", layout="wide")
st.title("Tapal DD 900gm Collectible – Daily Survey Updater")
st.caption("Upload the raw survey CSV plus the two master workbooks. "
           "Get back the two updated workbooks.")

c1, c2, c3 = st.columns(3)
survey_file = c1.file_uploader("1. Raw survey CSV", type=["csv"],
                               help="Tapal_Survey_Question_... .csv")
cs_file = c2.file_uploader("2. Collectible Summary master", type=["xlsx"])
dd_file = c3.file_uploader("3. DD 900gm CP master", type=["xlsx"])

ready = survey_file is not None and cs_file is not None and dd_file is not None
if st.button("Update both files", type="primary", disabled=not ready):
    with st.spinner("Matching stores, appending days and recalculating…"):
        try:
            st.session_state["result"] = run_update(
                survey_file.getvalue(), cs_file.getvalue(), dd_file.getvalue())
        except Exception as exc:  # show a readable message instead of a stack trace
            st.session_state.pop("result", None)
            st.error(f"Could not process the files: {exc}")

res = st.session_state.get("result")
if res:
    last = max(res["dates"])
    stamp = last.strftime("%d%b%Y")
    cs_info, dd_info = res["cs_info"], res["dd_info"]
    matched_cs = sum(len(s["matched_keys"]) for s in cs_info["sheets"])
    st.success(f"Done. Added {', '.join(d.strftime('%d-%b') for d in res['dates'])} "
               f"to both files.")

    m1, m2, m3, m4 = st.columns(4)
    m1.metric("Survey submissions", res["submissions"])
    m2.metric("Stores surveyed", res["stores_surveyed"])
    m3.metric("Stores matched in master", matched_cs)
    m4.metric("Survey stores not found", len(cs_info["unmatched"]))

    d1, d2 = st.columns(2)
    d1.download_button("⬇ Download Collectible Summary (updated)",
                       data=res["cs_bytes"],
                       file_name=f"Collectible_Summary_BA_Stores_{stamp}.xlsx",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    d2.download_button("⬇ Download DD 900gm CP Survey (updated)",
                       data=res["dd_bytes"],
                       file_name=f"DD_900gm_Mug_CP_Survey_{stamp}.xlsx",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    t1, t2, t3, t4 = st.tabs(["Daily totals", "Store report", "Duplicates", "Needs review"])
    with t1:
        st.dataframe(pd.DataFrame(
            [{"Date": d.strftime("%d-%b-%Y"), "Units sold (survey total)": v}
             for d, v in sorted(res["daily_totals"].items())]),
            hide_index=True, use_container_width=True)
    with t2:
        st.dataframe(pd.DataFrame(res["report"]), hide_index=True, use_container_width=True)
    with t3:
        if res["duplicates"]:
            st.caption("Stores with more than one submission on the same day. The latest "
                       "submission that has answers was used.")
            st.dataframe(pd.DataFrame(res["duplicates"]), hide_index=True, use_container_width=True)
        else:
            st.info("No duplicate submissions.")
    with t4:
        if cs_info["fuzzy"] or dd_info["fuzzy"]:
            st.warning("Store names matched by similarity (please verify):")
            st.dataframe(pd.DataFrame(
                [{"Survey name": s, "Master name": w} for s, w in cs_info["fuzzy"]]),
                hide_index=True, use_container_width=True)
        if cs_info["unmatched"]:
            st.error("These survey stores are not in the master files and were NOT updated. "
                     "Add them to the master first, then re-run:")
            st.write(cs_info["unmatched"])
        if not cs_info["fuzzy"] and not cs_info["unmatched"]:
            st.info("All survey stores matched.")

with st.expander("How the update works"):
    st.markdown(
        "- **Sell-out:** Q2 (Units Sold) is written into a new day column for each store. "
        "Blank answers are left empty.\n"
        "- **Stock status:** Q1 (CP stock) above 0 → *Available*; 0, *Not Available* or "
        "*Out Of Stock* → *Not Available*. The status header moves to the latest day.\n"
        "- **Remarks:** *Stock finished* is added when a store goes from Available to "
        "Not Available (or answers Out Of Stock), and removed when it becomes Available again.\n"
        "- **Duplicates:** if a store has several submissions on one day, the latest one "
        "with answers is used.\n"
        "- **Existing history** (earlier days, totals, summary blocks) is not changed.\n"
        "- Each output has an **Update Log** sheet listing what changed."
    )